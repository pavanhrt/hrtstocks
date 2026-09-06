from __future__ import annotations

import hashlib
import json
import re
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "tmp" / "source_extracts"
OUT.mkdir(parents=True, exist_ok=True)


def digest(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def safe_name(path: Path) -> str:
    rel = str(path.relative_to(ROOT))
    return re.sub(r"[^A-Za-z0-9._-]+", "_", rel)


def extract_pdf(path: Path) -> dict:
    reader = PdfReader(str(path))
    chunks = []
    page_lengths = []
    for i, page in enumerate(reader.pages, 1):
        try:
            text = page.extract_text() or ""
        except Exception as exc:
            text = f"[TEXT EXTRACTION ERROR: {exc}]"
        page_lengths.append(len(text.strip()))
        chunks.append(f"\n===== PAGE {i} =====\n{text}")
    out = OUT / f"{safe_name(path)}.txt"
    out.write_text("".join(chunks), encoding="utf-8")
    return {
        "type": "pdf",
        "pages": len(reader.pages),
        "page_text_lengths": page_lengths,
        "text_output": str(out.relative_to(ROOT)),
    }


def extract_docx(path: Path) -> dict:
    doc = Document(str(path))
    chunks = []
    for p in doc.paragraphs:
        if p.text.strip():
            chunks.append(f"[P:{p.style.name}] {p.text}")
    for ti, table in enumerate(doc.tables, 1):
        chunks.append(f"\n[TABLE {ti}]")
        for row in table.rows:
            chunks.append(" | ".join(cell.text.replace("\n", " / ") for cell in row.cells))
    out = OUT / f"{safe_name(path)}.txt"
    out.write_text("\n".join(chunks), encoding="utf-8")
    return {
        "type": "docx",
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "inline_shapes": len(doc.inline_shapes),
        "text_chars": sum(len(x) for x in chunks),
        "text_output": str(out.relative_to(ROOT)),
    }


def extract_xlsx(path: Path) -> dict:
    wb_formula = load_workbook(path, data_only=False, read_only=False)
    wb_values = load_workbook(path, data_only=True, read_only=False)
    chunks = []
    sheets = []
    for ws in wb_formula.worksheets:
        vws = wb_values[ws.title]
        sheets.append({
            "name": ws.title,
            "state": ws.sheet_state,
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "merged_ranges": [str(r) for r in ws.merged_cells.ranges],
        })
        chunks.append(f"\n===== SHEET {ws.title} ({ws.sheet_state}) {ws.max_row}x{ws.max_column} =====")
        for row in ws.iter_rows():
            vals = []
            for cell in row:
                if cell.value is None:
                    continue
                cached = vws[cell.coordinate].value
                value = str(cell.value).replace("\n", " / ")
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    value += f" [cached={cached!r}]"
                vals.append(f"{cell.coordinate}={value}")
            if vals:
                chunks.append(" | ".join(vals))
    out = OUT / f"{safe_name(path)}.txt"
    out.write_text("\n".join(chunks), encoding="utf-8")
    return {"type": "xlsx", "sheets": sheets, "text_output": str(out.relative_to(ROOT))}


def extract_html(path: Path) -> dict:
    text = path.read_text(encoding="utf-8", errors="replace")
    out = OUT / f"{safe_name(path)}.txt"
    out.write_text(text, encoding="utf-8")
    return {"type": "html", "chars": len(text), "text_output": str(out.relative_to(ROOT))}


def extract_zip(path: Path) -> dict:
    with zipfile.ZipFile(path) as zf:
        return {"type": "zip", "entries": [{"name": z.filename, "bytes": z.file_size} for z in zf.infolist()]}


files = sorted(p for p in ROOT.rglob("*") if p.is_file() and ".git" not in p.parts and "source_extracts" not in p.parts and p.name != Path(__file__).name)
manifest = []
seen = {}
for path in files:
    sha = digest(path)
    record = {
        "path": str(path.relative_to(ROOT)),
        "bytes": path.stat().st_size,
        "sha256": sha,
        "duplicate_of": seen.get(sha),
    }
    if sha not in seen:
        seen[sha] = record["path"]
        try:
            ext = path.suffix.lower()
            if ext == ".pdf":
                record.update(extract_pdf(path))
            elif ext == ".docx":
                record.update(extract_docx(path))
            elif ext == ".xlsx":
                record.update(extract_xlsx(path))
            elif ext in {".html", ".htm"}:
                record.update(extract_html(path))
            elif ext == ".zip":
                record.update(extract_zip(path))
            else:
                record["type"] = ext.lstrip(".") or "unknown"
        except Exception as exc:
            record["error"] = repr(exc)
    manifest.append(record)

(OUT / "manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
print(json.dumps(manifest, indent=2))
